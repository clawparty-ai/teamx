#!/usr/bin/env python3
# Build an EDITABLE 16:9 PPTX for the teamx FDE pitch.
# Same shallow/admin light theme as the HTML pages, but every text is a real
# text frame so the CEO/team can edit inline.
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

# ---- theme (teamx admin light) ----
BG      = RGBColor(0xF8, 0xF8, 0xF8)
CARD    = RGBColor(0xFF, 0xFF, 0xFF)
FG      = RGBColor(0x33, 0x33, 0x33)
MUT     = RGBColor(0x90, 0x93, 0x99)
PRIMARY = RGBColor(0x4A, 0x90, 0xE5)
PRIM_H  = RGBColor(0x35, 0x7A, 0xBD)
PRIM_L  = RGBColor(0xEB, 0xF3, 0xFC)
BORDER  = RGBColor(0xEB, 0xEE, 0xF5)
GREEN   = RGBColor(0x67, 0xC2, 0x3A)
FONT = "Microsoft YaHei"

EMU_IN = 914400
W, H = 13.333, 7.5
PAD = 0.45

def _set_font(run, size, bold=False, color=FG):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    # ensure east-asian font follows
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', FONT)

def _fill_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG

def _add_rect(slide, x, y, w, h, fill=None, line=None, round_=False, radius=0.06):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try: sp.adjustments[0] = radius
        except Exception: pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(0.75)
    sp.shadow.inherit = False
    return sp

def _add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
              space_after=4, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph = list of (text,size,bold,color)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for (t, size, bold, color) in para:
            r = p.add_run(); r.text = t; _set_font(r, size, bold, color)
    return tb

def _tag(slide, x, y, text, color=PRIMARY):
    sp = _add_rect(slide, x, y, 2.4, 0.30, fill=PRIM_L, line=RGBColor(0xCF, 0xE2, 0xFB), round_=True, radius=0.5)
    tf = sp.text_frame; tf.word_wrap = False
    tf.margin_left = Emu(90000); tf.margin_right = Emu(90000); tf.margin_top=0; tf.margin_bottom=0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = text; _set_font(r, 10.5, True, PRIMARY)
    return sp

def _header(slide, tagtext):
    # logo box
    lx = _add_rect(slide, PAD, PAD, 0.42, 0.42, fill=PRIMARY, round_=True, radius=0.25)
    tf = lx.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "Tx"; _set_font(r, 15, True, RGBColor(0xFF,0xFF,0xFF))
    _add_text(slide, PAD+0.52, PAD+0.03, 1.4, 0.4, [[("teamx", 16, True, PRIMARY)]], anchor=MSO_ANCHOR.MIDDLE)
    _tag(slide, W-PAD-4.4, PAD+0.06, tagtext)

def _foot(slide, left, page):
    _add_text(slide, PAD, H-PAD-0.25, 8, 0.3, [[(left, 10, False, MUT)]])
    _add_text(slide, W-PAD-1.2, H-PAD-0.25, 1.2, 0.3, [[(page, 10, False, MUT)]], align=PP_ALIGN.RIGHT)

def _card(slide, x, y, w, h):
    _add_rect(slide, x, y, w, h, fill=CARD, line=BORDER, round_=True, radius=0.05)
    return x, y, w, h

def _card_title(slide, x, y, w, h, icon, title, body, body_color=MUT, body_size=11.5, title_size=13):
    # icon line
    _add_text(slide, x+0.16, y+0.12, w-0.3, 0.3, [[(icon, 15, False, FG)]])
    _add_text(slide, x+0.16, y+0.42, w-0.32, 0.3, [[(title, title_size, True, FG)]])
    _add_text(slide, x+0.16, y+0.70, w-0.32, h-0.8, [[(body, body_size, False, body_color)]], line_spacing=1.12)

def _lead(slide, x, y, w, h, text_runs):
    _add_rect(slide, x, y, w, h, fill=PRIM_L, line=RGBColor(0xCF,0xE2,0xFB), round_=True, radius=0.06)
    tb = slide.shapes.add_textbox(Inches(x+0.2), Inches(y+0.12), Inches(w-0.4), Inches(h-0.24))
    tf = tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for para in text_runs:
        p = tf.paragraphs[0] if tf.paragraphs[0].runs==[] and para is text_runs[0] else tf.add_paragraph()
        p.line_spacing=1.1
        for (t, size, bold, color) in para:
            r=p.add_run(); r.text=t; _set_font(r, size, bold, color)

def _check_list(slide, x, y, w, items, size=11.5, gap=0.24):
    lines = []
    for it in items:
        lines.append([("✓  ", size, True, GREEN), (it, size, False, FG)])
    _add_text(slide, x, y, w, 2.0, lines, space_after=gap, line_spacing=1.0)

# ---------------------------------------------------------------- helpers end

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

# ============ P1 overview ============
s = prs.slides.add_slide(BLANK); _fill_bg(s)
_header(s, "FDE · 组织与治理底座")

_add_text(s, PAD, 1.0, 9.2, 1.6, [
    [("把 AI 工程师，像一支训练有素的远程团队，派驻到业务现场", 24, True, FG)],
])
_add_text(s, PAD, 1.72, 9.2, 0.7, [[("teamx 是 AI FDE 的组织与治理底座", 17, True, PRIMARY)]])
_add_text(s, PAD, 2.14, 11.6, 0.75, [[(
    "让“人类做方向、AI 工程师做执行”规模化运转——任务分派、状态流转、人工验收、全程留痕、成本配额，一个工作台全部看清。",
    13, False, MUT)]], line_spacing=1.2)

# three cards
cw, ch, gap = 3.70, 2.1, 0.28
_x = PAD
for (ic, ti, bo) in [
    ("🎯", "对内：统一指挥", "所有 FDE 项目进同一套团队操作系统：谁在做什么、卡在哪、花了多少，一目了然。"),
    ("🤝", "对外：可信交付", "客户开一个团队就看到“自己的 AI 工程师”与实时工作台，按任务验收，不是黑盒 API。"),
    ("🛡️", "全程留痕可控", "每步进账本、交付可追溯、人工 Approve 写闸门 + 成本配额，让 AI 进生产能过风控。"),
]:
    _card(s, _x, 3.15, cw, ch); _card_title(s, _x, 3.15, cw, ch, ic, ti, bo)
    _x += cw + gap

# second row: two slim cards
cw2, ch2, gap2 = 5.71, 1.0, 0.28
_x = PAD
for (num, lbl, sub) in [
    ("7×24", "常驻推进", "AI FDE 前向部署在客户工作流里，休息时间也推进，不占人头、随叫随到。"),
    ("按任务 · 配额度", "成本看得见", "成本按任务结算，内置配额护栏；客户放权不再担心“无限烧钱”。"),
]:
    _card(s, _x, 5.42, cw2, ch2)
    _add_text(s, _x+0.16, 5.52, 2.3, 0.3, [[(num, 14, True, PRIMARY)]])
    _add_text(s, _x+2.4, 5.52, 1.2, 0.3, [[(lbl, 12, True, FG)]])
    _add_text(s, _x+3.4, 5.50, cw2-3.5, 0.9, [[(sub, 10.5, False, MUT)]], line_spacing=1.1)
    _x += cw2 + gap2

_lead(s, PAD, 6.6, W-2*PAD, 0.52,
      [[("一句话　", 12, True, PRIMARY),
        ("把 AI 工程师像一支训练有素的远程团队派驻到客户业务里——客户只管派目标与验收，过程透明、成本可控、随时可交接。", 12, True, FG)]])
_foot(s, "teamx · 共享目标团队协作 / AI FDE 运行治理底座", "01 · 一页概览")

# ============ P2 why ============
s = prs.slides.add_slide(BLANK); _fill_bg(s)
_header(s, "DETAIL 01 · 问题与解法")
_add_text(s, PAD, 0.98, 10, 0.7, [[("AI FDE 规模化之前，先回答客户的", 22, True, FG)],
                                  [("四个不放心", 22, True, PRIMARY)]])

# table-ish via 4 rows of rectangles
rows = [
    ("失控", "一堆 Agent 同时在客户那边跑，没有“谁在干什么”的统一视图。",
     "一个团队工作台：成员/任务/状态/目标全在账本上，随时看全局。"),
    ("不可验收", "干完了，但交付物没有标准、没有过程证据，客户不敢签收。",
     "任务状态机 + 交付物 + 人工 Approve 闸门：验收通过才算完成，全程留痕。"),
    ("成本黑洞", "Agent 无限探索烧 token，客户不敢放权给 AI。",
     "每成员/每任务成本配额护栏，超限即停；成本按任务透明结算。"),
    ("不可交接", "人一走、Agent 一关，上下文与知识全丢。",
     "契约化 TEAM.md + 引用带原文快照 + 审计账本：换人换 Agent 知识仍在。"),
]
y0, rh, rgap = 1.85, 0.78, 0.10
c1, c2, c3 = 2.0, 4.8, 5.1
for i,(a,b,c) in enumerate(rows):
    y = y0 + i*(rh+rgap)
    _add_rect(s, PAD, y, c1, rh, fill=CARD, line=BORDER, round_=True)
    _add_rect(s, PAD+c1+0.08, y, c2, rh, fill=CARD, line=BORDER, round_=True)
    _add_rect(s, PAD+c1+0.08+c2+0.08, y, c3, rh, fill=PRIM_L, line=RGBColor(0xCF,0xE2,0xFB), round_=True)
    _add_text(s, PAD+0.12, y+0.06, c1-0.24, rh-0.1, [[(a, 12.5, True, FG)]], anchor=MSO_ANCHOR.MIDDLE)
    _add_text(s, PAD+c1+0.20, y+0.04, c2-0.3, rh-0.06, [[(b, 10.5, False, MUT)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    _add_text(s, PAD+c1+0.16+c2+0.08, y+0.04, c3-0.3, rh-0.06, [[(c, 10.5, False, FG)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)

# bottom three cards
cw, ch, gap = 3.70, 1.6, 0.28
_x = PAD
for (ic, ti, bo) in [
    ("🧩", "不是又一个写代码的 Agent", "写代码的 Agent 很多；teamx 是让“人类 owner + 一群 AI FDE”组成可信团队的那一层。"),
    ("🔁", "人机混合，不是替代", "专家 + AI FDE 同一团队协作，AI 让每位工程师产能翻倍，内部推行阻力小。"),
    ("📐", "试点变方法论", "一次成功 FDE 项目沉淀成团队模板，新客户开箱即用，交付可复制。"),
]:
    _card(s, _x, 5.35, cw, ch); _card_title(s, _x, 5.35, cw, ch, ic, ti, bo, title_size=13, body_size=11)
    _x += cw + gap
_foot(s, "为什么 AI FDE 需要 teamx", "02 · 详细介绍")

# ============ P3 internal ============
s = prs.slides.add_slide(BLANK); _fill_bg(s)
_header(s, "DETAIL 02 · 面向企业内部")
_add_text(s, PAD, 0.98, 11, 0.7, [[("在公司内部推行 FDE，你不只缺 Agent，", 22, True, FG)],
                                  [("更缺秩序", 22, True, PRIMARY)]])

cards = [
    ("🗂️", "统一指挥层（CEO / Lead）", ["所有 FDE 项目进同一套“团队操作系统”",
        "每支 AI 团队：目标 / 任务 / 进度 / 成本一块屏看清",
        "从几个试点到全公司方法论，有统一底座"]),
    ("📜", "方法论可复制", ["TEAM.md 契约：角色 / 流程 / 验收标准写进团队定义",
        "成功项目沉淀为模板，新项目/客户开箱即用",
        "跨团队复用 SOP，不依赖某个人或某个 Agent"]),
    ("🛡️", "合规与风控闸门", ["人工 Approve 写闸门：AI 产出先审后入",
        "全程 ledger 留痕，可导出审计",
        "让 AI 进内部系统过得了安全与法务"]),
    ("🧑‍🤝‍🧑", "人机混合团队", ["专家与 AI FDE 同一团队，不是替代",
        "AI 承接重复/监控/初稿，专家做判断与终审",
        "每位工程师产能翻倍，推行阻力最小化"]),
]
cw, ch, gap = 5.90, 1.95, 0.24
for i,(ic, ti, items) in enumerate(cards):
    x = PAD + (i%2)*(cw+gap)
    y = 1.78 + (i//2)*(ch+gap)
    _card(s, x, y, cw, ch)
    _add_text(s, x+0.16, y+0.10, 0.5, 0.3, [[(ic, 16, False, FG)]])
    _add_text(s, x+0.52, y+0.12, cw-0.6, 0.3, [[(ti, 13.5, True, FG)]])
    _check_list(s, x+0.18, y+0.50, cw-0.35, items, size=11, gap=0.16)

# flow lead
_lead(s, PAD, 5.95, W-2*PAD, 0.72, [
    [("典型内部项目流　", 12, True, PRIMARY),
     ("建团队 / 定义 TEAM.md　→　AI FDE 认领任务　→　产出进入验收队列　→　专家 Approve / 打回　→　沉淀为模板复用", 12, False, FG)]])
_foot(s, "面向企业内部推行 FDE", "03 · 详细介绍")

# ============ P4 customer ============
s = prs.slides.add_slide(BLANK); _fill_bg(s)
_header(s, "DETAIL 03 · 面向公司客户")
_add_text(s, PAD, 0.98, 11, 1.0, [[("对客户，不卖“Agent”，卖一支看得见、管得住", 22, True, FG)],
                                  [("的 AI 工程师团队", 22, True, PRIMARY)]])

cards4 = [
    ("👁️", "客户看到的是“团队”，不是黑盒",
     "开一个团队就看到“自己的 AI 工程师”名单、实时工作台与任务进展——像雇了一支远程团队。"),
    ("🧾", "按任务验收，不是按人头付费",
     "派活 → 状态机流转 → 交付物 → 客户验收/打回；成本按任务结算，自带配额护栏。"),
    ("🔍", "过程可审计，交付可追溯",
     "每步进账本，引用带原文快照——解决“AI 不可信”最大顾虑，签收有依据。"),
    ("🌐", "7×24 前向部署",
     "AI FDE 常驻客户工作流，休息时间也推进；客户只做方向与验收。"),
]
cw, ch, gap = 5.90, 2.15, 0.24
for i,(ic, ti, bo) in enumerate(cards4):
    x = PAD + (i%2)*(cw+gap)
    y = 2.0 + (i//2)*(ch+gap)
    _card(s, x, y, cw, ch); _card_title(s, x, y, cw, ch, ic, ti, bo, title_size=13.5, body_size=11.5)

# chips + flow
_lead(s, PAD, 6.5, W-2*PAD, 0.5,
      [[("客户体验的一周　", 11.5, True, PRIMARY),
        ("开团队/看到 FDE 名单　→　派第一个目标　→　7×24 推进并汇报　→　客户验收交付物　→　成本明细清晰可查", 11.5, False, FG)]])
_foot(s, "teamx · FDE 交付底座", "04 · 详细介绍")

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "teamx-FDE-pitch-editable.pptx")
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides._sldIdLst))
